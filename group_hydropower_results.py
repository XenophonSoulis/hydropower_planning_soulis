import openpyxl

address=r'C:\Works\paper_hydropower_XKS\edits_hydropower_ped'
up_fluctuation=2 # Set it to a negative value to group everything into one group; set it to 0 to prevent grouping
down_fluctuation=5 # Set it to a negative value to group everything into one group; set it to 0 to prevent grouping
cont_fluctuation=10 # Set it to a negative value to group everything into one group; set it to 0 to prevent grouping


def read_input():
    '''Reads the input 'config.txt' file'''
    with open(f'{address}\\config.txt') as config:
        int_read_config=lambda:int(config.readline()[:-1].split(sep=':')[1])
        float_read_config=lambda:float(config.readline()[:-1].split(sep=':')[1])
        bool_read_config=lambda:bool(int(config.readline()[:-1].split(sep=':')[1]))
        str_read_config=lambda:config.readline()[:-1].split(sep=':')[1]
        float_read_config()
        return float_read_config()
    
cell_size=read_input()

# Square fluctuations to compare with squared distances
if up_fluctuation>0:
    up_fluctuation**=2
if down_fluctuation>0:
    down_fluctuation**=2
if cont_fluctuation>0:
    cont_fluctuation**=2

def get_value(x):
    return x.value

iexcel=openpyxl.load_workbook(f'{address}\\results.xlsx')
isheet=iexcel.active
iiter=iter(isheet.rows)
titles=list(map(get_value,next(iiter)))

def dist(p,q):
    return (p[0]-q[0])**2+(p[1]-q[1])**2+(p[2]-q[2])**2

def fix(up,down,cont):
    return ((up[0]/cell_size,up[1]/cell_size,up[2]/cell_size),
            (down[0]/cell_size,down[1]/cell_size,down[2]/cell_size),
            (cont[0]/cell_size,cont[1]/cell_size,cont[2]/cell_size))

def check_dist(p,q,fluctuation):
    if fluctuation>0:
        return dist(p,q)<fluctuation
    elif fluctuation==0:
        return False
    else:
        return True

def in_groups(up,down,cont):
    up,down,cont=fix(up,down,cont)
    use_groups=[]
    for i,(up0,down0,cont0) in enumerate(fixed_coords):
        if check_dist(up,up0,up_fluctuation) and check_dist(down,down0,down_fluctuation) and check_dist(cont,cont0,cont_fluctuation):
            use_groups.append(i)
    return use_groups

groups=[]
fixed_coords=[]
all_results=[]

for raw_row in iiter:
    row=list(map(get_value,raw_row))
    all_results.append(row)
    up=row[1:4]
    down=row[4:7]
    cont=*row[7:9],row[3]
    use_groups=in_groups(up,down,cont)
    if use_groups:
        for i in use_groups:
            groups[i].append(row)
    else:
        groups.append([row])
        fixed_coords.append(fix(up,down,cont))

main_page=[]

for i,group in enumerate(groups):
    main_page.append((*group[0],f'Group {i}'))

oexcel=openpyxl.Workbook()

omain=oexcel.active
omain.title='Main'
omain.append((*titles,'Group sheet'))
for row in main_page:
    omain.append(row)

for i,group in enumerate(groups):
    sheet=oexcel.create_sheet(f'Group {i}')
    sheet.append(titles)
    for row in group:
        sheet.append(row)

ototal=oexcel.create_sheet(f'Full list')
ototal.append(titles)
for row in all_results:
    ototal.append(row)

oexcel.save(f'{address}\\grouped_results.xlsx')
